#!/usr/bin/env python3
"""Run isolated XLSX-to-CSV benchmarks and emit a machine-readable report."""

from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.metadata
import io
import json
import os
import platform
import random
import subprocess
import sys
import time
import ctypes
from datetime import date, datetime, time as datetime_time
from pathlib import Path
from statistics import median

try:
    import resource
except ImportError:  # Windows does not provide the Unix resource module.
    resource = None  # type: ignore[assignment]


ENGINES = ("turboxl", "calamine", "openpyxl")
def format_number(value: float | int | bool) -> str:
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if value != value:
        return "#NUM!"
    if value == float("inf"):
        return "#DIV/0!"
    if value == float("-inf"):
        return "-#DIV/0!"
    if value == int(value) and abs(value) < 1e15:
        return str(int(value))
    rendered = f"{value:.6f}".rstrip("0").rstrip(".")
    return rendered or "0"


def normalize_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, (int, float, bool)):
        return format_number(value)
    if isinstance(value, datetime):
        return value.replace(microsecond=0).isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, datetime_time):
        return value.replace(microsecond=0).isoformat()
    return str(value)


def write_rows(rows: object, output: io.StringIO) -> int:
    writer = csv.writer(output, lineterminator="\n")
    row_count = 0
    for row in rows:
        writer.writerow([normalize_cell(value) for value in row])
        row_count += 1
    return row_count


def run_engine(engine: str, xlsx: str, sheet_index: int) -> tuple[str, int]:
    if engine == "turboxl":
        import turboxl

        rendered = turboxl.read_sheet_to_csv(xlsx, sheet_index)
        return rendered, rendered.count("\n")

    output = io.StringIO()
    if engine == "calamine":
        import python_calamine

        workbook = python_calamine.load_workbook(xlsx)
        sheet = workbook.get_sheet_by_index(sheet_index)
        rows = sheet.iter_rows()
    elif engine == "openpyxl":
        import openpyxl

        workbook = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
        sheet = workbook.worksheets[sheet_index]
        rows = sheet.iter_rows(values_only=True)
    else:
        raise ValueError(f"Unknown engine: {engine}")

    row_count = write_rows(rows, output)
    return output.getvalue(), row_count


def peak_rss_kb() -> int:
    if resource is not None:
        peak_rss = int(resource.getrusage(resource.RUSAGE_SELF).ru_maxrss)
        if sys.platform == "darwin":
            peak_rss //= 1024
        return peak_rss

    if os.name == "nt":
        from ctypes import wintypes

        class ProcessMemoryCounters(ctypes.Structure):
            _fields_ = [
                ("cb", wintypes.DWORD),
                ("PageFaultCount", wintypes.DWORD),
                ("PeakWorkingSetSize", ctypes.c_size_t),
                ("WorkingSetSize", ctypes.c_size_t),
                ("QuotaPeakPagedPoolUsage", ctypes.c_size_t),
                ("QuotaPagedPoolUsage", ctypes.c_size_t),
                ("QuotaPeakNonPagedPoolUsage", ctypes.c_size_t),
                ("QuotaNonPagedPoolUsage", ctypes.c_size_t),
                ("PagefileUsage", ctypes.c_size_t),
                ("PeakPagefileUsage", ctypes.c_size_t),
            ]

        counters = ProcessMemoryCounters()
        counters.cb = ctypes.sizeof(counters)
        get_current_process = ctypes.windll.kernel32.GetCurrentProcess
        get_current_process.argtypes = []
        get_current_process.restype = wintypes.HANDLE
        get_process_memory_info = ctypes.windll.psapi.GetProcessMemoryInfo
        get_process_memory_info.argtypes = [
            wintypes.HANDLE,
            ctypes.POINTER(ProcessMemoryCounters),
            wintypes.DWORD,
        ]
        get_process_memory_info.restype = wintypes.BOOL
        handle = get_current_process()
        succeeded = get_process_memory_info(
            handle, ctypes.byref(counters), counters.cb
        )
        if not succeeded:
            raise ctypes.WinError()
        return int(counters.PeakWorkingSetSize // 1024)

    raise RuntimeError(f"Peak RSS measurement is unsupported on {sys.platform}")


def worker(args: argparse.Namespace) -> int:
    rendered, row_count = run_engine(args.engine, args.xlsx, args.sheet_index)
    payload = {
        "engine": args.engine,
        "bytes": len(rendered.encode("utf-8")),
        "rows": row_count,
        "sha256": hashlib.sha256(rendered.encode("utf-8")).hexdigest(),
        "peak_rss_kb": peak_rss_kb(),
    }
    print(json.dumps(payload, separators=(",", ":")))
    return 0


def package_versions() -> dict[str, str]:
    versions: dict[str, str] = {}
    for distribution in ("turboxl", "python-calamine", "openpyxl"):
        try:
            versions[distribution] = importlib.metadata.version(distribution)
        except importlib.metadata.PackageNotFoundError:
            versions[distribution] = "not-installed"
    return versions


def cpu_details() -> dict[str, str]:
    details = {
        "architecture": platform.machine(),
        "processor": platform.processor(),
        "logical_cpus": str(os.cpu_count() or "unknown"),
    }
    try:
        output = subprocess.check_output(
            ["lscpu"], text=True, stderr=subprocess.DEVNULL
        )
        wanted = {
            "Model name": "model_name",
            "Vendor ID": "vendor_id",
            "CPU(s)": "lscpu_count",
            "Thread(s) per core": "threads_per_core",
        }
        for line in output.splitlines():
            key, separator, value = line.partition(":")
            if separator and key.strip() in wanted:
                details[wanted[key.strip()]] = value.strip()
    except (FileNotFoundError, subprocess.CalledProcessError):
        pass
    if os.name == "nt":
        try:
            details["model_name"] = subprocess.check_output(
                [
                    "powershell.exe",
                    "-NoProfile",
                    "-Command",
                    "(Get-CimInstance Win32_Processor | "
                    "Select-Object -First 1 -ExpandProperty Name).Trim()",
                ],
                text=True,
                stderr=subprocess.DEVNULL,
            ).strip()
        except (FileNotFoundError, subprocess.CalledProcessError):
            pass
    return details


def timed_worker(
    script: Path, engine: str, xlsx: str, sheet_index: int
) -> dict[str, object]:
    command = [
        sys.executable,
        str(script),
        "--worker",
        "--engine",
        engine,
        "--xlsx",
        xlsx,
        "--sheet-index",
        str(sheet_index),
    ]
    started = time.perf_counter()
    completed = subprocess.run(command, capture_output=True, text=True)
    elapsed = time.perf_counter() - started
    if completed.returncode:
        raise RuntimeError(
            f"{engine} worker exited with {completed.returncode}: {completed.stderr.strip()}"
        )
    result = json.loads(completed.stdout)
    result["seconds"] = elapsed
    return result


def summarize(samples: list[dict[str, object]]) -> dict[str, object]:
    seconds = [float(sample["seconds"]) for sample in samples]
    rss = [int(sample["peak_rss_kb"]) for sample in samples]
    representative = samples[-1]
    return {
        "median_seconds": median(seconds),
        "min_seconds": min(seconds),
        "max_seconds": max(seconds),
        "median_peak_rss_mb": median(rss) / 1024,
        "rows": representative["rows"],
        "bytes": representative["bytes"],
        "sha256": representative["sha256"],
        "samples": samples,
    }


def orchestrator(args: argparse.Namespace) -> int:
    script = Path(__file__).resolve()
    xlsx = str(Path(args.xlsx).resolve())

    # Warm each implementation independently before collecting measurements.
    for engine in ENGINES:
        timed_worker(script, engine, xlsx, args.sheet_index)

    samples: dict[str, list[dict[str, object]]] = {engine: [] for engine in ENGINES}
    randomizer = random.Random(20260920)
    for _ in range(args.rounds):
        order = list(ENGINES)
        randomizer.shuffle(order)
        for engine in order:
            samples[engine].append(timed_worker(script, engine, xlsx, args.sheet_index))

    summaries = {engine: summarize(values) for engine, values in samples.items()}
    baseline = float(summaries["turboxl"]["median_seconds"])
    for engine, summary in summaries.items():
        summary["relative_to_turboxl"] = float(summary["median_seconds"]) / baseline

    hashes = {engine: summary["sha256"] for engine, summary in summaries.items()}
    report = {
        "schema_version": 1,
        "timestamp_utc": datetime.utcnow().replace(microsecond=0).isoformat() + "Z",
        "input": {
            "filename": Path(xlsx).name,
            "size_bytes": Path(xlsx).stat().st_size,
            "sheet_index": args.sheet_index,
            "rounds": args.rounds,
            "warmup_rounds": 1,
        },
        "machine": {
            "hostname": platform.node(),
            "platform": platform.platform(),
            "python": platform.python_version(),
            "cloud_provider": os.environ.get("BENCHMARK_CLOUD_PROVIDER", "unknown"),
            "cloud_region": os.environ.get("BENCHMARK_CLOUD_REGION", "unknown"),
            "cloud_vm_size": os.environ.get("BENCHMARK_CLOUD_VM_SIZE", "unknown"),
            "cpu": cpu_details(),
        },
        "packages": package_versions(),
        "results": summaries,
        "parity": {
            "all_hashes_match": len(set(hashes.values())) == 1,
            "hashes": hashes,
        },
    }
    json.dump(report, sys.stdout, indent=2, sort_keys=True)
    print()
    return 0


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--sheet-index", type=int, default=0)
    parser.add_argument("--rounds", type=int, default=7)
    parser.add_argument("--worker", action="store_true", help=argparse.SUPPRESS)
    parser.add_argument("--engine", choices=ENGINES, help=argparse.SUPPRESS)
    args = parser.parse_args()
    if args.rounds < 1:
        parser.error("--rounds must be at least 1")
    if args.worker and not args.engine:
        parser.error("--engine is required with --worker")
    return args


def main() -> int:
    args = parse_args()
    return worker(args) if args.worker else orchestrator(args)


if __name__ == "__main__":
    raise SystemExit(main())

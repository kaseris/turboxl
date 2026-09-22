"""Installed-wheel pandas adapter checks, including Windows file-handle cleanup."""

from __future__ import annotations

from datetime import datetime
import io
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest import mock

import openpyxl
import pandas as pd
from pandas.testing import assert_frame_equal
import turboxl
import turboxl.pandas as adapter

from fixtures import workbook


class PandasEngineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        source = Path(__file__).resolve().parents[2]
        assert source not in Path(turboxl.__file__).resolve().parents, turboxl.__file__

    def setUp(self):
        self.previous = pd.ExcelFile._engines.get("turboxl")
        pd.ExcelFile._engines.pop("turboxl", None)
        self.addCleanup(self.restore_registry)
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "sample.xlsx"
        workbook(self.path)

    def restore_registry(self):
        pd.ExcelFile._engines.pop("turboxl", None)
        if self.previous is not None:
            pd.ExcelFile._engines["turboxl"] = self.previous

    def register(self):
        self.assertIsNone(adapter.register())

    def assert_released(self, path):
        moved = path.with_suffix(".moved")
        path.rename(moved)
        moved.unlink()

    def make_simple(self):
        book = openpyxl.Workbook()
        try:
            sheet = book.active
            sheet.title = "Values"
            sheet.append(["name", "number", "flag", "when", "optional"])
            sheet.append(["alpha", 1, True, datetime(2024, 1, 2), None])
            sheet.append(["beta", 2, False, datetime(2024, 1, 3), "x"])
            hidden = book.create_sheet("Hidden")
            hidden.sheet_state = "hidden"
            hidden.append(["secret"])
            book.create_sheet("Empty")
            book.save(self.path)
        finally:
            book.close()

    def test_import_is_inert_and_registration_is_idempotent(self):
        result = subprocess.run(
            [sys.executable, "-c", "import sys, turboxl; assert 'pandas' not in sys.modules"],
            check=False, capture_output=True, text=True, cwd=self.temp.name,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.register()
        registered = pd.ExcelFile._engines["turboxl"]
        self.register()
        self.assertIs(pd.ExcelFile._engines["turboxl"], registered)

    def test_registration_errors_leave_registry_intact(self):
        pd.ExcelFile._engines["turboxl"] = object()
        with self.assertRaisesRegex(RuntimeError, "another provider"):
            adapter.register()
        existing = pd.ExcelFile._engines.pop("turboxl")
        with mock.patch.object(adapter.importlib, "import_module", side_effect=ImportError("missing")):
            with self.assertRaisesRegex(ImportError, "install turboxl"):
                adapter.register()
        with mock.patch.object(pd, "__version__", "3.1.0"):
            with self.assertRaisesRegex(RuntimeError, "unsupported"):
                adapter.register()
        with mock.patch.object(pd.ExcelFile, "_engines", None):
            with self.assertRaisesRegex(RuntimeError, "incompatible"):
                adapter.register()
        self.assertIsNotNone(existing)
        self.assertNotIn("turboxl", pd.ExcelFile._engines)

    def test_dataframe_parity_and_sheet_selection(self):
        self.make_simple()
        self.register()
        for kwargs in (
            {},
            {"header": None},
            {"skiprows": 1, "header": None},
            {"usecols": "A:C"},
            {"dtype": {"number": "float64"}},
            {"converters": {"number": lambda value: str(value)}},
            {"nrows": 1},
            {"parse_dates": ["when"]},
        ):
            with self.subTest(kwargs=kwargs):
                actual = pd.read_excel(self.path, engine="turboxl", **kwargs)
                for engine in ("calamine", "openpyxl"):
                    expected = pd.read_excel(self.path, engine=engine, **kwargs)
                    self.assertEqual(list(actual.dtypes), list(expected.dtypes))
                    assert_frame_equal(actual.isna(), expected.isna())
                    missing = object()
                    left = actual.astype(object).where(actual.notna(), missing)
                    right = expected.astype(object).where(expected.notna(), missing)
                    assert_frame_equal(left, right, check_dtype=True)
        self.assertEqual(
            pd.read_excel(self.path, engine="turboxl", sheet_name="Hidden", header=None).iloc[0, 0],
            "secret",
        )
        self.assertEqual(
            pd.read_excel(self.path, engine="turboxl", sheet_name=1, header=None).iloc[0, 0],
            "secret",
        )
        sheets = pd.read_excel(self.path, engine="turboxl", sheet_name=None)
        self.assertEqual(list(sheets), ["Values", "Hidden", "Empty"])
        self.assertTrue(sheets["Empty"].empty)
        self.assert_released(self.path)

    def test_sources_and_owned_streams(self):
        self.make_simple()
        self.register()
        archive = self.path.read_bytes()
        for source in (str(self.path), self.path, io.BytesIO(archive)):
            with self.subTest(source=type(source)):
                result = pd.read_excel(source, engine="turboxl")
                self.assertEqual(result["name"].tolist(), ["alpha", "beta"])
                if isinstance(source, io.BytesIO):
                    self.assertFalse(source.closed)
                    source.close()
        with self.path.open("rb") as source:
            result = pd.read_excel(source, engine="turboxl",
                                   engine_kwargs={"max_cells": 100})
            self.assertEqual(len(result), 2)
            self.assertFalse(source.closed)
            with self.assertRaises(RuntimeError):
                pd.read_excel(source, engine="turboxl",
                              engine_kwargs={"max_cells": 1})
            self.assertFalse(source.closed)
        self.assert_released(self.path)

    def test_reuse_close_and_retained_sheet(self):
        self.make_simple()
        self.register()
        with pd.ExcelFile(self.path, engine="turboxl") as excel:
            self.assertEqual(excel.sheet_names, ["Values", "Hidden", "Empty"])
            retained = excel.book.get_sheet_by_name("Values")
            self.assertEqual(excel.parse("Values")["name"].tolist(), ["alpha", "beta"])
            self.assertEqual(excel.parse("Hidden", header=None).iloc[0, 0], "secret")
        excel.close()
        with self.assertRaises(RuntimeError):
            retained.to_python()
        self.assert_released(self.path)

    def test_failure_paths_release_file(self):
        self.register()
        cases = (
            ("bad_archive", b"not a ZIP", {}),
            ("bad_sheet", None, {"sheet_name": "Missing"}),
            ("cell_limit", None, {"sheet_name": "Sparse", "engine_kwargs": {"max_cells": 10}}),
            ("converter", None, {"header": None, "converters": {0: lambda _: 1 / 0}}),
            ("bad_kwarg", None, {"engine_kwargs": {"unknown": 1}}),
        )
        for name, contents, kwargs in cases:
            with self.subTest(name=name):
                path = Path(self.temp.name) / f"{name}.xlsx"
                if contents is None:
                    workbook(path)
                else:
                    path.write_bytes(contents)
                with self.assertRaises(Exception):
                    pd.read_excel(path, engine="turboxl", **kwargs)
                self.assert_released(path)

    def test_sparse_limit_and_missing_values(self):
        self.register()
        sparse = pd.read_excel(
            self.path, engine="turboxl", sheet_name="Sparse", header=None, nrows=5
        )
        self.assertEqual(sparse.shape, (5, 4))
        self.assertEqual(sparse.iloc[4, 3], "gap")
        scalar_path = Path(self.temp.name) / "scalars.xlsx"
        workbook(scalar_path, scalars=True)
        with turboxl.load_workbook(scalar_path) as native:
            cells = native.get_sheet_by_name("Data").to_python()[0]
        self.assertIsNone(cells[0])
        self.assertIsNone(cells[5])  # Excel error and blank are intentionally indistinguishable.
        result = pd.read_excel(
            scalar_path, engine="turboxl", header=None, keep_default_na=False
        )
        self.assertTrue(pd.isna(result.iloc[0, 0]))
        self.assertTrue(pd.isna(result.iloc[0, 5]))
        self.assert_released(scalar_path)
        self.assert_released(self.path)


if __name__ == "__main__":
    unittest.main()
